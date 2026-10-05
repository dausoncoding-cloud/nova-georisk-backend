"""
Report generation — Doc 0's reporting/export step: summary PDF and
multi-sheet Excel workbooks covering FIRRIS index results and model
validation metrics.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime, timezone

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


@dataclass
class IndexSummary:
    """One FIRRIS index's summary stats, for the report header block."""

    name: str  # e.g. "Flood Hazard Index (H)"
    mean: float
    min: float
    max: float
    classification: str  # dominant/representative class label


@dataclass
class ReportContext:
    project_name: str
    aoi_name: str
    generated_at: datetime = field(default_factory=lambda: datetime.now(timezone.utc))
    index_summaries: list[IndexSummary] = field(default_factory=list)
    regression_metrics: dict[str, float] | None = None
    classification_metrics: dict[str, float] | None = None
    class_area_statistics: list[dict] = field(default_factory=list)
    notes: str | None = None


def generate_pdf_report(context: ReportContext, output_path: str) -> str:
    """
    Build a summary PDF: title block, FIRRIS index summary table, and
    (if provided) validation metric tables. Returns the path written.
    """
    doc = SimpleDocTemplate(output_path, pagesize=A4, topMargin=2 * cm, bottomMargin=2 * cm)
    styles = getSampleStyleSheet()
    story = []

    story.append(Paragraph("NOVA-GeoRisk Intelligence Suite — FIRRIS Report", styles["Title"]))
    story.append(Spacer(1, 0.3 * cm))
    story.append(Paragraph(f"Project: {context.project_name}", styles["Normal"]))
    story.append(Paragraph(f"AOI: {context.aoi_name}", styles["Normal"]))
    story.append(Paragraph(f"Generated: {context.generated_at.strftime('%Y-%m-%d %H:%M UTC')}", styles["Normal"]))
    story.append(Spacer(1, 0.6 * cm))

    if context.index_summaries:
        story.append(Paragraph("Flood Risk Index Summary", styles["Heading2"]))
        table_data = [["Index", "Mean", "Min", "Max", "Classification"]]
        for s in context.index_summaries:
            table_data.append([s.name, f"{s.mean:.3f}", f"{s.min:.3f}", f"{s.max:.3f}", s.classification])
        story.append(_styled_table(table_data))
        story.append(Spacer(1, 0.6 * cm))

    if context.regression_metrics:
        story.append(Paragraph("Regression Validation", styles["Heading2"]))
        table_data = [["Metric", "Value"]] + [[k.upper(), f"{v:.4f}"] for k, v in context.regression_metrics.items()]
        story.append(_styled_table(table_data))
        story.append(Spacer(1, 0.6 * cm))

    if context.classification_metrics:
        story.append(Paragraph("Classification Validation", styles["Heading2"]))
        table_data = [["Metric", "Value"]] + [
            [k.replace("_", " ").title(), f"{v:.4f}"] for k, v in context.classification_metrics.items()
        ]
        story.append(_styled_table(table_data))
        story.append(Spacer(1, 0.6 * cm))

    if context.class_area_statistics:
        story.append(Paragraph("Delivered Raster Area by Class", styles["Heading2"]))
        story.append(Paragraph("Percentages use valid delivered cells as the denominator; no class is an absolute safety threshold.", styles["Normal"]))
        rows = [["Product", "Class", "Cells", "Area (ha)", "Valid area (%)"]]
        for item in context.class_area_statistics:
            rows.append([str(item["product"]), str(item["class_value"]), str(item["cells"]),
                         f"{item['area_ha']:.3f}", f"{item['percent_of_valid']:.2f}"])
        story.append(_styled_table(rows))
        story.append(Spacer(1, 0.6 * cm))

    if context.notes:
        story.append(Paragraph("Notes", styles["Heading2"]))
        story.append(Paragraph(context.notes, styles["Normal"]))

    doc.build(story)
    return output_path


def _styled_table(data: list[list[str]]) -> Table:
    table = Table(data, hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F2F2")]),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    return table


def generate_excel_report(sheets: dict[str, pd.DataFrame], output_path: str) -> str:
    """
    Write a multi-sheet Excel workbook — one sheet per named DataFrame
    (e.g. {"Hazard Index": df, "Validation Metrics": df}). Header row
    is styled; column widths are auto-sized to content.
    """
    if not sheets:
        raise ValueError("At least one sheet is required.")

    wb = Workbook()
    wb.remove(wb.active)  # drop the default blank sheet

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    for sheet_name, df in sheets.items():
        ws = wb.create_sheet(title=sheet_name[:31])  # Excel sheet-name length limit

        ws.append(list(df.columns))
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font

        for row in df.itertuples(index=False):
            ws.append(list(row))

        for col_idx, column in enumerate(df.columns, start=1):
            max_len = max([len(str(column))] + [len(str(v)) for v in df[column]])
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 40)

    wb.save(output_path)
    return output_path


def export_csv(df: pd.DataFrame, output_path: str) -> str:
    df.to_csv(output_path, index=False)
    return output_path
