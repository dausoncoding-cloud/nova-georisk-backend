import tempfile
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from app.services.reporting.report_builder import (
    IndexSummary,
    ReportContext,
    export_csv,
    generate_excel_report,
    generate_pdf_report,
)


def _sample_context():
    return ReportContext(
        project_name="Dar es Salaam Pilot",
        aoi_name="Msimbazi Basin",
        index_summaries=[
            IndexSummary(name="Flood Hazard Index (H)", mean=0.62, min=0.10, max=0.95, classification="High"),
            IndexSummary(name="Flood Risk Index (FRI)", mean=0.48, min=0.05, max=0.88, classification="Moderate"),
        ],
        regression_metrics={"rmse": 1.23, "r_squared": 0.87, "mae": 0.95},
        classification_metrics={"overall_accuracy": 91.5, "cohens_kappa": 0.82},
        class_area_statistics=[{"product": "Flood Extent", "class_value": 1,
                                "cells": 42, "area_ha": 12.5, "percent_of_valid": 25.0}],
        notes="Pilot run over the Msimbazi River basin, wet season 2026.",
    )


def test_generate_pdf_report_creates_valid_pdf():
    context = _sample_context()
    with tempfile.TemporaryDirectory() as tmpdir:
        path = str(Path(tmpdir) / "report.pdf")
        generate_pdf_report(context, path)

        assert Path(path).exists()
        reader = PdfReader(path)
        assert len(reader.pages) >= 1

        full_text = "".join(page.extract_text() or "" for page in reader.pages)
        assert "Dar es Salaam Pilot" in full_text
        assert "Msimbazi Basin" in full_text
        assert "Flood Hazard Index" in full_text
        assert "RMSE" in full_text.upper() or "rmse" in full_text.lower()
        assert "Delivered Raster Area by Class" in full_text
        assert "Flood Extent" in full_text
        assert "12.500" in full_text


def test_generate_pdf_report_without_optional_sections():
    """A minimal context (no metrics, no notes) should still produce a valid PDF."""
    context = ReportContext(project_name="Minimal", aoi_name="Test AOI")
    with tempfile.TemporaryDirectory() as tmpdir:
        path = str(Path(tmpdir) / "minimal.pdf")
        generate_pdf_report(context, path)
        reader = PdfReader(path)
        assert len(reader.pages) >= 1


def test_generate_excel_report_creates_multi_sheet_workbook():
    sheets = {
        "Hazard Index": pd.DataFrame({"pixel_id": [1, 2, 3], "hazard_score": [0.2, 0.5, 0.9]}),
        "Validation Metrics": pd.DataFrame({"metric": ["RMSE", "R2"], "value": [1.23, 0.87]}),
    }
    with tempfile.TemporaryDirectory() as tmpdir:
        path = str(Path(tmpdir) / "report.xlsx")
        generate_excel_report(sheets, path)

        wb = load_workbook(path)
        assert set(wb.sheetnames) == {"Hazard Index", "Validation Metrics"}

        ws = wb["Hazard Index"]
        assert [c.value for c in ws[1]] == ["pixel_id", "hazard_score"]
        assert ws[2][0].value == 1
        assert ws[2][1].value == 0.2


def test_generate_excel_report_truncates_long_sheet_names():
    long_name = "A" * 50
    sheets = {long_name: pd.DataFrame({"x": [1]})}
    with tempfile.TemporaryDirectory() as tmpdir:
        path = str(Path(tmpdir) / "long_name.xlsx")
        generate_excel_report(sheets, path)
        wb = load_workbook(path)
        assert len(wb.sheetnames[0]) <= 31


def test_generate_excel_report_rejects_empty_sheets():
    with tempfile.TemporaryDirectory() as tmpdir:
        with pytest.raises(ValueError):
            generate_excel_report({}, str(Path(tmpdir) / "empty.xlsx"))


def test_export_csv_round_trip():
    df = pd.DataFrame({"a": [1, 2], "b": ["x", "y"]})
    with tempfile.TemporaryDirectory() as tmpdir:
        path = str(Path(tmpdir) / "out.csv")
        export_csv(df, path)
        read_back = pd.read_csv(path)
        pd.testing.assert_frame_equal(read_back, df)
