"""Synthetic delivery/audit fixtures cover U27–U33, not scientific truth."""
import json
import uuid
import zipfile
from copy import deepcopy
from pathlib import Path
from types import SimpleNamespace
from xml.etree import ElementTree as ET

import geopandas as gpd
import numpy as np
import pandas as pd
import pytest
import rasterio
from geoalchemy2.shape import from_shape
from shapely.geometry import Polygon, MultiPolygon, Point, LineString, mapping
from rasterio.transform import from_origin
from openpyxl import load_workbook
from pydantic import ValidationError
from pypdf import PdfReader

from app.core.config import get_settings, Settings
from app.platform.result_exports import artifact_entry, build_result_exports
from app.schemas.result_delivery import ResultAnalytics, ResultInterpretation
from app.services.maps.export import RasterSpec, write_cog
from app.services.maps.vector_formats import export_vector_formats, KML
from app.services.reporting.evidence import quantitative_delivery, interpret_evidence, display_bounds, verified_artifact
from app.services.reporting.complete_report import complete_report, W
from app.services.tasks.execution import initialize_execution, validate_execution, append_event, verify_events, execution_record, validate_workload, redact_snapshot


def _raster(directory, values, *, key="synthetic_classes", dtype="uint8", crs=3857):
    spec = RasterSpec(from_origin(4000000, -100000, 10, 10), crs)
    path = directory / (key + ".tif")
    write_cog(str(path), np.asarray(values, dtype=dtype), spec, nodata=255 if dtype == "uint8" else -9999.)
    return artifact_entry(path, label=key, media_type="image/tiff", artifact_type="raster", result_version=1,
        role="product", product_key=key, format_name="cog", gis_metadata={"crs": f"EPSG:{crs}", "units": "synthetic class" if dtype == "uint8" else "synthetic score", "nodata": 255 if dtype == "uint8" else -9999., "legend": [{"value": 0, "label": "dry", "color": "#FFFFFF"}, {"value": 1, "label": "wet", "color": "#0000FF"}, {"value": 2, "label": "absent", "color": "#00FF00"}]})


def test_all_class_areas_include_zero_classes_units_denominator_and_nodata(tmp_path):
    entry = _raster(tmp_path, [[0, 1], [1, 255]])
    data = quantitative_delivery(tmp_path, {"classes": entry}, {}, {})
    assert len(data["class_areas"]) == 3
    rows = {item["class_value"]: item for item in data["class_areas"]}
    assert rows[0]["cells"] == 1 and rows[0]["area_m2"] == 100
    assert rows[1]["area_ha"] == .02 and rows[1]["area_km2"] == .0002
    assert rows[1]["percent_of_valid"] == pytest.approx(200 / 3)
    assert rows[2]["cells"] == rows[2]["area_m2"] == 0
    assert all(item["denominator_area_m2"] == 300 for item in rows.values())
    assert sum(item["percent_of_valid"] for item in rows.values()) == pytest.approx(100)
    assert data["series"][0]["kind"] == "class_area"


def test_continuous_scores_get_descriptive_histograms_without_invented_classes(tmp_path):
    entry = _raster(tmp_path, [[.1, .5], [.8, -9999]], dtype="float32")
    data = quantitative_delivery(tmp_path, {"scores": entry}, {}, {})
    assert not data["class_areas"]
    series = data["series"][0]
    assert series["kind"] == "histogram" and sum(point["count"] for point in series["points"]) == 3
    assert "display-only" in series["basis"]


def test_geographic_area_and_projected_display_bounds_are_explicit(tmp_path):
    path = tmp_path / "geographic.tif"
    write_cog(str(path), np.ones((2, 2), dtype="uint8"), RasterSpec(from_origin(36, -1, .01, .01), 4326), nodata=255)
    entry = artifact_entry(path, label="synthetic", media_type="image/tiff", artifact_type="raster", result_version=1,
        role="product", product_key="flood_extent", format_name="cog", gis_metadata={"crs": "EPSG:4326", "nodata": 255})
    data = quantitative_delivery(tmp_path, {"extent": entry}, {}, {})
    assert data["class_areas"][1]["area_method"] == "ellipsoidal_geodesic"
    bounds = display_bounds("EPSG:3857", dict(west=4000000, south=-100000, east=4000010, north=-99990))
    assert 35 < bounds["west"] < bounds["east"] < 37 and -1 < bounds["south"] < bounds["north"] < 0
    with pytest.raises(ValueError):
        display_bounds("EPSG:4326", dict(west=170, south=-5, east=190, north=5))


@pytest.mark.parametrize("mutation", ["checksum", "crs", "nodata", "invalid_values"])
def test_quantitative_delivery_rejects_inconsistent_artifacts(tmp_path, mutation):
    entry = _raster(tmp_path, [[0, 1], [1, 0]])
    if mutation == "checksum":
        (tmp_path / entry["path"]).write_bytes(b"synthetic tamper")
    elif mutation == "invalid_values":
        with rasterio.open(tmp_path / entry["path"], "r+", IGNORE_COG_LAYOUT_BREAK="YES") as raster:
            raster.write(np.full((2, 2), 255, dtype="uint8"), 1)
        from app.platform.result_exports import artifact_fingerprint
        entry["file_size_bytes"], entry["checksum_sha256"] = artifact_fingerprint(tmp_path / entry["path"])
    else:
        entry["gis_metadata"][mutation] = "EPSG:4326" if mutation == "crs" else -9999
    with pytest.raises(ValueError):
        quantitative_delivery(tmp_path, {"source": entry}, {}, {})


def _temporal():
    summary = {"change_statistics": {"before_inundated_area_m2": 0., "after_inundated_area_m2": 100., "net_inundated_change_m2": 100.}}
    provenance = {"comparison": {"before_timestamp": "2020-06-01T00:00:00+00:00", "after_timestamp": "2020-06-02T00:00:00+00:00", "observation_definition": "synthetic fixed definition"}}
    return summary, provenance


def test_temporal_series_uses_only_comparable_observed_instants_and_qualifies_change(tmp_path):
    summary, provenance = _temporal()
    data = quantitative_delivery(tmp_path, {}, summary, provenance)
    series = data["series"][0]
    assert series["kind"] == "observed_time_series" and series["units"] == "m2"
    assert [point["value"] for point in series["points"]] == [0, 100]
    text = interpret_evidence(data, summary, provenance)
    assert text["independent_scientific_validation"] is False
    assert any("undefined" in item["text"] and "do not establish cause" in item["text"] for item in text["findings"])
    assert all(item["evidence_refs"] for item in text["findings"] + text["recommendations"])
    assert not quantitative_delivery(tmp_path, {}, {"created_at": "2020-01-01", "version": 2}, {})["series"]


@pytest.mark.parametrize("mutation", ["definition", "order", "invalid_area"])
def test_temporal_delivery_fails_closed_without_comparable_evidence(tmp_path, mutation):
    summary, provenance = _temporal()
    if mutation == "definition":
        provenance["comparison"].pop("observation_definition")
    elif mutation == "order":
        provenance["comparison"]["after_timestamp"] = provenance["comparison"]["before_timestamp"]
    else:
        summary["change_statistics"]["before_inundated_area_m2"] = -1
    with pytest.raises(ValueError):
        quantitative_delivery(tmp_path, {}, summary, provenance)


def _vector(directory, geometry=None, properties=None, *, features=None):
    features = features if features is not None else [{"type": "Feature", "geometry": mapping(geometry or Point(36, -1)), "properties": properties or {"synthetic_long_attribute_name": "synthetic<&>", "original_number": 12345678901234}}]
    path = directory / "source.geojson"
    path.write_text(json.dumps({"type": "FeatureCollection", "features": features}), encoding="utf-8")
    return artifact_entry(path, label="synthetic", media_type="application/geo+json", artifact_type="vector", result_version=1,
        role="product", product_key="synthetic_vector", format_name="geojson", gis_metadata={"crs": "EPSG:4326", "units": "synthetic source geometry"})


@pytest.mark.parametrize("geometry", [Point(36, -1), LineString([(36, -1), (36.5, -1.5)]),
    Polygon([(36, -2), (37, -2), (37, -1), (36, -1)], holes=[[(36.2, -1.8), (36.4, -1.8), (36.4, -1.4), (36.2, -1.4)]]),
    MultiPolygon([Polygon([(36, -2), (36.2, -2), (36.2, -1.8), (36, -1.8)]), Polygon([(36.4, -2), (36.6, -2), (36.6, -1.8), (36.4, -1.8)])])])
def test_vector_formats_preserve_crs_geometry_and_typed_attribute_sidecar(tmp_path, geometry):
    entry = _vector(tmp_path, geometry)
    exports = export_vector_formats(tmp_path, "synthetic", entry, {"source": "synthetic fixture"})
    gpkg = gpd.read_file(exports["gpkg"][0])
    assert gpkg.crs.to_epsg() == 4326 and gpkg.geometry[0].equals(geometry)
    assert gpkg["original_number"][0] == 12345678901234
    root = ET.parse(exports["kml"][0])
    assert root.find(".//{" + KML + "}value").text == "synthetic<&>"
    if geometry.geom_type == "Polygon" and geometry.interiors:
        assert root.find(".//{" + KML + "}innerBoundaryIs") is not None
    with zipfile.ZipFile(exports["shapefile_zip"][0]) as archive:
        assert {"features.shp", "features.shx", "features.dbf", "features.prj", "attributes.json"} <= set(archive.namelist())
        metadata = json.loads(archive.read("attributes.json"))
        assert metadata["properties"][0]["original_number"] == 12345678901234
        assert "synthetic_long_attribute_name" in metadata["field_map"].values()


def test_empty_observed_extent_has_empty_valid_vector_exports(tmp_path):
    entry = _vector(tmp_path, features=[])
    exports = export_vector_formats(tmp_path, "empty_extent", entry, {})
    assert len(gpd.read_file(exports["gpkg"][0])) == 0
    assert ET.parse(exports["kml"][0]).find(".//{" + KML + "}Placemark") is None


def test_vector_export_rejects_unknown_crs_and_partitions_mixed_geometry_without_substitution(tmp_path):
    entry = _vector(tmp_path)
    entry["gis_metadata"]["crs"] = "EPSG:3857"
    with pytest.raises(ValueError, match="WGS84"):
        export_vector_formats(tmp_path, "unsupported", entry, {})
    features = [{"type": "Feature", "geometry": mapping(geometry), "properties": {}} for geometry in (Point(36, -1), Polygon([(36, -2), (37, -2), (37, -1), (36, -1)]))]
    entry = _vector(tmp_path, features=features)
    exports = export_vector_formats(tmp_path, "mixed", entry, {})
    assert len(gpd.read_file(exports["gpkg"][0])) == 2
    with zipfile.ZipFile(exports["shapefile_zip"][0]) as archive:
        sidecar = json.loads(archive.read("attributes.json"))
        assert sidecar["layers"] == {"point": {"geometry_family": "Point", "source_feature_indices": [0]},
                                    "polygon": {"geometry_family": "Polygon", "source_feature_indices": [1]}}
        archive.extractall(tmp_path / "synthetic_mixed_layers")
    for name, geometry in zip(("point", "polygon"), (Point(36, -1), Polygon([(36, -2), (37, -2), (37, -1), (36, -1)]))):
        assert gpd.read_file(tmp_path / "synthetic_mixed_layers" / (name + ".shp")).geometry[0].equals(geometry)
    assert len(list(ET.parse(exports["kml"][0]).getroot().iter("{" + KML + "}Placemark"))) == 2


def test_vector_export_rejects_unsupported_geometry_collection(tmp_path):
    from shapely.geometry import GeometryCollection
    entry = _vector(tmp_path, features=[{"type": "Feature", "geometry": mapping(GeometryCollection([Point(36, -1)])), "properties": {}}])
    with pytest.raises(ValueError, match="invalid/unsupported"):
        export_vector_formats(tmp_path, "unsupported_geometry", entry, {})


def test_complete_reports_keep_evidence_limits_and_escape_markup_and_formulas(tmp_path):
    summary, provenance = _temporal()
    provenance["synthetic_untrusted_source_name"] = "=HYPERLINK(\"unsafe\")"
    data = quantitative_delivery(tmp_path, {}, summary, provenance)
    interpretation = interpret_evidence(data, summary, provenance)
    reports = complete_report(tmp_path, {"project_id": "synthetic<&>", "result_version": 1}, {"crs": "EPSG:3857", "units": "m2"}, summary, provenance, data, interpretation)
    text = " ".join(page.extract_text() for page in PdfReader(reports["complete_report_pdf"][0]).pages)
    assert "synthetic<&>" in text and "Source approval" in text and "Validation and its scope" in text
    workbook = load_workbook(reports["complete_report_excel"][0])
    assert {"Class areas", "Provenance", "Validation", "Interpretation"} <= set(workbook.sheetnames)
    assert not any(cell.data_type == "f" for sheet in workbook for row in sheet for cell in row)
    with zipfile.ZipFile(reports["complete_report_word"][0]) as archive:
        assert {"[Content_Types].xml", "_rels/.rels", "word/document.xml"} <= set(archive.namelist())
        document = ET.fromstring(archive.read("word/document.xml"))
        word_text = " ".join(node.text or "" for node in document.iter("{" + W + "}t"))
        assert "synthetic<&>" in word_text and "independent scientific validation" in word_text
    assert "Observed net inundated-area change" in reports["complete_report_csv"][0].read_text()


def test_shared_exports_include_new_reports_vectors_and_package_content(tmp_path):
    entry = _raster(tmp_path, [[0, 1], [1, 255]])
    vector = _vector(tmp_path)
    summary = {"status": "completed"}
    exports = build_result_exports(tmp_path, task_id=str(uuid.uuid4()), project_id=str(uuid.uuid4()), aoi_id=str(uuid.uuid4()),
        engine_key="firris", engine_version="1.0", result_type="synthetic", result_version=1,
        summary=summary, provenance={"source": "synthetic fixture"}, gis_metadata={"crs": "EPSG:3857"}, product_entries={"classes": entry, "source": vector})
    assert summary["delivery"]["class_areas"] and summary["interpretation"]["findings"]
    for name in ("complete_report_word", "quantitative_data", "evidence_interpretation", "source_gpkg", "source_kml", "source_shapefile_zip"):
        assert exports[name]["role"] == "export" and exports[name]["checksum_sha256"]
    with zipfile.ZipFile(tmp_path / exports["report_package"]["path"]) as package:
        assert "reports/complete-evidence-report.docx" in package.namelist()
        assert "reports/source.gpkg" in package.namelist()


def _task_fixture():
    aoi = SimpleNamespace(geometry=from_shape(MultiPolygon([Polygon([(36, -2), (37, -2), (37, -1), (36, -1)])]), srid=4326), source_lineage={"source": "synthetic"})
    task = SimpleNamespace(input_params={"parameters": {"synthetic": [0, 1]}}, result_payload=None, progress_pct=0)
    initialize_execution(task, aoi)
    return task, aoi


def test_execution_records_exact_inputs_aoi_events_cache_policy_and_retry_link():
    task, aoi = _task_fixture()
    validate_execution(task, aoi)
    task.progress_pct = 50
    append_event(task, "progress")
    record = execution_record(task)
    assert record["aoi_snapshot"]["type"] == "MultiPolygon"
    assert record["aoi_source_lineage"]["source"] == "synthetic"
    assert record["origin"] == "legacy_worker" and record["events"][0]["actor_id"] is None
    assert record["cache_policy"] == "revalidate_registered_sources_no_automatic_result_or_gee_reuse"
    verify_events(record)
    retry = uuid.uuid4()
    initialize_execution(task, aoi, context=SimpleNamespace(user_id=retry), retry_of=retry)
    assert execution_record(task)["retry_of"] == str(retry)
    assert execution_record(task)["events"][0]["actor_id"] == str(retry)


@pytest.mark.parametrize("mutation", ["parameters", "aoi", "code", "event"])
def test_execution_fails_closed_on_modified_parameters_aoi_code_or_audit_event(mutation):
    task, aoi = _task_fixture()
    if mutation == "parameters":
        task.input_params["parameters"]["synthetic"] = [1, 0]
    elif mutation == "aoi":
        aoi.geometry = from_shape(MultiPolygon([Polygon([(35, -2), (37, -2), (37, -1), (35, -1)])]), srid=4326)
    elif mutation == "code":
        task.result_payload["execution_record"]["implementation_sha256"] = "different"
    else:
        task.result_payload["execution_record"]["events"][0]["progress_pct"] = 100
    with pytest.raises(ValueError):
        validate_execution(task, aoi)


def test_capacity_limits_and_export_secret_redaction_are_explicit(monkeypatch):
    settings = get_settings()
    monkeypatch.setattr(settings, "firris_max_prepared_cells", 3)
    with pytest.raises(ValueError, match="capacity"):
        validate_workload({"parameters": {"workflow": {"label_layer": [[0, 1], [1, 0]], "feature_layers": {"synthetic": [[0, 1], [1, 0]]}}}})
    monkeypatch.setattr(settings, "firris_max_request_bytes", 32)
    with pytest.raises(ValueError, match="capacity"):
        validate_workload({"synthetic": "x" * 100})
    with pytest.raises(ValueError):
        Settings(firris_max_pending_tasks_per_organization=0)
    assert redact_snapshot({"parameters": {"private_key": "synthetic-secret", "value": 1}}) == {"parameters": {"private_key": "[redacted]", "value": 1}}


def test_existing_probability_display_bands_supply_qualified_areas_without_changing_raw_scores(tmp_path):
    entry = _raster(tmp_path, [[0., .3], [.5, .9]], key="flood_probability", dtype="float32")
    entry["gis_metadata"]["legend"] = {"entries": [{"label": "0.00–0.20", "color": "#FFFFFF"}, {"label": "0.20–0.40", "color": "#AAAAAA"}, {"label": "0.40–0.60", "color": "#777777"}, {"label": "0.60–1.00", "color": "#000000"}]}
    original = (tmp_path / entry["path"]).read_bytes()
    data = quantitative_delivery(tmp_path, {"score_cog": entry}, {}, {})
    assert len(data["class_areas"]) == 4 and all(row["percent_of_valid"] == 25 for row in data["class_areas"])
    assert all("display legend" in row["classification_basis"] for row in data["class_areas"])
    assert (tmp_path / entry["path"]).read_bytes() == original
    assert any(series["kind"] == "histogram" for series in data["series"])


def test_long_dbf_attributes_are_declared_references_with_full_sidecar_values(tmp_path):
    value = "synthetic α" * 100
    entry = _vector(tmp_path, properties={"long_attribute": value})
    exports = export_vector_formats(tmp_path, "long_attributes", entry, {})
    with zipfile.ZipFile(exports["shapefile_zip"][0]) as archive:
        sidecar = json.loads(archive.read("attributes.json"))
        assert sidecar["properties"][0]["long_attribute"] == value
        archive.extractall(tmp_path / "read_synthetic")
    frame = gpd.read_file(tmp_path / "read_synthetic" / "features.shp")
    assert frame["f0000000"][0].startswith("[full value: attributes.json")


def test_parent_timeout_callback_records_only_actual_hard_failure(monkeypatch):
    from app.workers import celery_tasks as module
    calls = []
    monkeypatch.setattr(module.Request, 'on_timeout', lambda self, soft, timeout: None)
    monkeypatch.setattr(module, 'persist_worker_failure', lambda task_id, error: calls.append((task_id, str(error))))
    request = object.__new__(module.FIRRISRequest)
    request._args = ['synthetic-task-id']
    request.on_timeout(True, 10)
    assert calls == []
    request.on_timeout(False, 20)
    assert calls == [('synthetic-task-id', 'Worker hard time limit exceeded')]
    assert module.run_engine_analysis.Request == 'app.workers.celery_tasks:FIRRISRequest'
    assert module.run_flood_screening_atlas.Request == 'app.workers.celery_tasks:FIRRISRequest'


def test_parent_worker_loss_callback_records_failure_without_fabricating_retry(monkeypatch):
    from app.workers import celery_tasks as module
    from billiard.exceptions import WorkerLostError
    calls = []
    monkeypatch.setattr(module.Request, 'on_failure', lambda *args: None)
    monkeypatch.setattr(module, 'persist_worker_failure', lambda task_id, error: calls.append((task_id, type(error))))
    request = object.__new__(module.FIRRISRequest)
    request._args = ['synthetic-task-id']
    request.on_failure(SimpleNamespace(exception=module.Retry()))
    assert calls == []
    request.on_failure(SimpleNamespace(exception=WorkerLostError('synthetic child loss')))
    assert calls == [('synthetic-task-id', WorkerLostError)]


def test_bundle3_contracts_and_both_exports_match_runtime():
    from app.main import app
    from app.bff.main import app as browser
    root = Path(__file__).resolve().parents[2]
    for runtime, filename in ((app, 'nova-internal-api.json'), (browser, 'nova-browser-api.json')):
        schema = runtime.openapi()
        assert json.loads((root / 'openapi' / filename).read_text()) == schema
        assert schema['paths']['/api/v1/tasks/{task_id}/archive']['get']['responses']['200']['content']['application/json']['schema']['$ref'].endswith('TaskArchiveResponse')
        models = schema['components']['schemas']
        assert {'analytics', 'interpretation'} <= set(models['ResultResponse']['properties'])
        assert {'display_bounds_wgs84', 'temporal_metadata'} <= set(models['ResultLayerResponse']['properties'])
        assert 'execution' in models['TaskStatusResponse']['properties']
        assert models['ResultInterpretation']['properties']['method']['const'] == 'deterministic_evidence_rules_v1'


@pytest.mark.parametrize('field', ['geometry', 'FID'])
def test_vector_exports_reject_reserved_attribute_overwrite(tmp_path, field):
    entry = _vector(tmp_path, properties={field: 'synthetic original attribute'})
    with pytest.raises(ValueError, match='reserved'):
        export_vector_formats(tmp_path, 'invalid_reserved_field', entry, {})


@pytest.mark.parametrize('value', [float('nan'), float('inf')])
def test_vector_exports_reject_nonfinite_scientific_attributes(tmp_path, value):
    entry = _vector(tmp_path, properties={'synthetic_invalid_observation': value})
    with pytest.raises(ValueError):
        export_vector_formats(tmp_path, 'invalid_nonfinite_value', entry, {})
